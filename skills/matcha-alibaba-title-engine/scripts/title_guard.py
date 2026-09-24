#!/usr/bin/env python3
"""Validate Alibaba matcha titles against length, casing, and rolling duplication rules."""

from __future__ import annotations

import argparse
import json
import re
import sys
from difflib import SequenceMatcher
from pathlib import Path


SMALL_WORDS = {
    "the", "a", "an", "and", "but", "for", "nor", "or", "so", "yet", "if",
    "at", "by", "in", "of", "on", "to", "up", "with",
}
ACRONYMS = {"OEM", "ODM", "B2B", "DIY", "BPA", "UV"}
PROCUREMENT_STARTS = (
    "custom ", "customized ", "personalized ", "oem ", "odm ",
    "factory ", "supplier ", "logo custom ", "private label ", "wholesale ",
)
BANNED_STOCK_TERMS = ("us stock", "usa stock", "in stock")
CORE_TERMS = (
    "matcha set", "matcha tea set", "matcha kit", "complete matcha kit",
    "matcha bowl", "matcha bowl set", "matcha whisk", "matcha travel set",
    "chasen", "chawan", "berry bowl",
)


def normalize_text(text: str) -> str:
    return " ".join(text.lower().split())


def tokenize(text: str) -> list[str]:
    return re.findall(r"[a-z0-9]+(?:[-'][a-z0-9]+)?", text.lower())


def longest_common_run(left: str, right: str) -> int:
    a, b = tokenize(left), tokenize(right)
    previous = [0] * (len(b) + 1)
    best = 0
    for token_a in a:
        current = [0]
        for j, token_b in enumerate(b, start=1):
            value = previous[j - 1] + 1 if token_a == token_b else 0
            current.append(value)
            best = max(best, value)
        previous = current
    return best


def load_ledger_titles(path: Path) -> list[str]:
    if not path.exists():
        return []
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read the XLSX ledger") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    if "标题台账" not in workbook.sheetnames:
        raise ValueError(f"Missing 标题台账 sheet in {path}")
    sheet = workbook["标题台账"]
    header_row = None
    title_col = None
    status_col = None
    for row_index, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        for col_index, value in enumerate(row, start=1):
            header = str(value).strip()
            if header == "英文标题":
                header_row, title_col = row_index, col_index
            elif header in {"版本状态", "状态"}:
                status_col = col_index
        if title_col:
            break
    if not title_col or not header_row:
        raise ValueError(f"Missing 英文标题 column in {path}")
    titles = []
    max_col = max(title_col, status_col or title_col)
    for row in sheet.iter_rows(min_row=header_row + 1, min_col=1, max_col=max_col, values_only=True):
        title = row[title_col - 1]
        status = row[status_col - 1] if status_col else None
        normalized_status = normalize_text(str(status)) if status is not None else ""
        if normalized_status in {"superseded", "inactive", "已替换", "停用", "失效"}:
            continue
        if title and str(title).strip():
            titles.append(str(title).strip())
    return titles


def load_recent_json(path: Path | None) -> list[str]:
    if not path or not path.exists():
        return []
    data = json.loads(path.read_text(encoding="utf-8"))
    records = data.get("titles", data) if isinstance(data, dict) else data
    titles = []
    for record in records:
        if isinstance(record, str):
            titles.append(record)
        elif isinstance(record, dict) and record.get("active", True) and record.get("title"):
            titles.append(str(record["title"]).strip())
    return titles


def title_case_issues(title: str) -> list[str]:
    raw_words = title.split()
    issues = []
    for index, raw in enumerate(raw_words):
        word = re.sub(r"^[^A-Za-z0-9]+|[^A-Za-z0-9]+$", "", raw)
        if not word or word.isdigit() or re.fullmatch(r"\d+(?:\.\d+)?(?:ml|oz|cm|in)?", word, re.I):
            continue
        if word.upper() in ACRONYMS:
            if word != word.upper():
                issues.append(f"{word} should be uppercase")
            continue
        lower = word.lower()
        if index > 0 and lower in SMALL_WORDS:
            if word != lower:
                issues.append(f"{word} should be lowercase")
        elif word[0].isalpha() and not word[0].isupper():
            issues.append(f"{word} should start uppercase")
    return issues


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--title", required=True)
    parser.add_argument("--ledger", type=Path)
    parser.add_argument("--recent-json", type=Path)
    parser.add_argument("--recent-title", action="append", default=[])
    parser.add_argument("--piece-count", type=int)
    parser.add_argument("--primary-phrase")
    parser.add_argument("--max-chars", type=int, default=128)
    parser.add_argument("--target-min", type=int, default=105)
    parser.add_argument("--target-max", type=int, default=125)
    parser.add_argument("--similarity-threshold", type=float, default=0.60)
    args = parser.parse_args()

    title = " ".join(args.title.split())
    ledger_titles = load_ledger_titles(args.ledger) if args.ledger else []
    recent_titles = load_recent_json(args.recent_json) + args.recent_title
    comparisons = (ledger_titles + recent_titles)[-5:]

    matches = []
    for prior in comparisons:
        matches.append({
            "title": prior,
            "similarity": round(SequenceMatcher(None, normalize_text(title), normalize_text(prior)).ratio(), 4),
            "longest_contiguous_words": longest_common_run(title, prior),
        })
    top_match = max(matches, key=lambda item: item["similarity"], default=None)
    longest_run = max((item["longest_contiguous_words"] for item in matches), default=0)

    failures = []
    warnings = []
    char_count = len(title)
    if char_count > args.max_chars:
        failures.append(f"Title is {char_count} characters; maximum is {args.max_chars}")
    elif not args.target_min <= char_count <= args.target_max:
        warnings.append(f"Title is outside the preferred {args.target_min}-{args.target_max} character range")

    invalid = sorted(set(re.findall(r"[@!?！？]", title)))
    if invalid:
        failures.append("Disallowed punctuation: " + " ".join(invalid))
    lowered = title.lower()
    if any(term in lowered for term in BANNED_STOCK_TERMS):
        failures.append("Stock wording conflicts with the current customization-first project rule")
    if not lowered.startswith(PROCUREMENT_STARTS):
        failures.append("Title must begin with a truthful procurement/customization expression")
    if not any(term in lowered for term in CORE_TERMS):
        failures.append("No recognized matcha/adjacent core product phrase found")

    primary_phrase = " ".join(args.primary_phrase.split()) if args.primary_phrase else None
    primary_protected = bool(primary_phrase and normalize_text(primary_phrase) in normalize_text(title))
    if primary_phrase and not primary_protected:
        failures.append(f"Protected primary phrase is missing: {primary_phrase}")

    casing = title_case_issues(title)
    if casing:
        failures.extend(casing)

    piece_matches = [int(value) for value in re.findall(r"\b(\d+)\s*(?:piece|pieces|pcs)\b", title, re.I)]
    if piece_matches:
        if args.piece_count is None:
            warnings.append("Piece count appears in title but no verified --piece-count was supplied")
        elif any(value != args.piece_count for value in piece_matches):
            failures.append(f"Title piece count does not match verified count {args.piece_count}")

    if top_match and top_match["similarity"] >= args.similarity_threshold:
        failures.append(f"Similarity {top_match['similarity']:.1%} reaches the {args.similarity_threshold:.0%} rewrite threshold")
    if longest_run >= 6:
        failures.append(f"Found a repeated contiguous sequence of {longest_run} words")

    result = {
        "title": title,
        "character_count": char_count,
        "comparison_count": len(comparisons),
        "maximum_similarity": top_match["similarity"] if top_match else 0,
        "most_similar_title": top_match["title"] if top_match else None,
        "longest_contiguous_words": longest_run,
        "primary_phrase": primary_phrase,
        "primary_phrase_protected": primary_protected if primary_phrase else None,
        "failures": failures,
        "warnings": warnings,
        "status": "pass" if not failures else "fail",
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if not failures else 1


if __name__ == "__main__":
    sys.exit(main())
