#!/usr/bin/env python3
"""Inspect and append immutable XLS/XLSX/CSV source versions to a skill registry."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path


KNOWN_HEADERS = {
    "关键词", "搜索指数", "搜索涨幅", "点击率", "卖家规模指数",
    "产品ID", "产品名称", "搜索曝光次数", "搜索点击次数", "搜索点击率", "询盘个数",
    "编号", "日期", "中文品名", "英文标题", "采购词", "核心词", "新增覆盖词",
    "真实属性词", "字符数（公式）", "A级权重", "最近5条最高相似度", "连续6词检查", "记录说明",
    "标题编号", "标题日期", "词级", "关键词/词组", "类别", "覆盖说明", "权重", "证据状态",
    "累计次数", "当日次数", "7日次数", "30日次数", "当日状态", "7日状态", "30日状态",
    "项目", "说明", "来源", "数据日期", "时间范围", "关键词/标题", "指标类型", "曝光", "点击",
    "关键词指数", "卖家指数", "证据等级", "备注/URL",
    "产品ID/SKU", "类目路径", "标题版本", "模型版本", "核心词保护（公式）", "版本状态", "上线日期",
    "主核心词", "次要入口词", "发布日期", "统计开始", "统计结束", "搜索曝光", "搜索点击", "询盘", "提交订单",
    "搜索CTR（公式）", "每千曝光询盘（公式）", "每千曝光订单（公式）", "观察天数（公式）",
    "价格变更", "主图变更", "广告变更", "可售/履约变更", "可比状态（公式）", "样本置信度（公式）",
    "学习状态（公式）", "Source文件/证据ID", "数据校验（公式）", "窗口重叠（公式）", "参数", "当前值", "用途", "更新规则",
    "候选标题", "产品相关", "流量分0-100", "转化分0-100", "证据分0-100", "产品差异化分0-5",
    "覆盖奖励分0-5", "硬门槛（公式）", "综合分（公式）", "证据ID/说明", "生效日期", "关键词/词组",
    "旧权重", "旧权重0-1", "原始建议变动", "单轮上限（公式）", "应用变动（公式）", "新权重（公式）", "新权重0-1（公式）", "合格窗口数", "证据ID", "调整原因", "状态",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def clean(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def detect_header(rows: list[list[object]]) -> tuple[int | None, list[str]]:
    best_index, best_headers, best_score = None, [], 0
    for index, row in enumerate(rows[:20], start=1):
        headers = [clean(value) for value in row]
        score = sum(1 for value in headers if value in KNOWN_HEADERS)
        if score > best_score:
            best_index, best_headers, best_score = index, headers, score
    return best_index, [value for value in best_headers if value]


def classify(sheet_names: list[str], headers: list[str]) -> str:
    header_set = set(headers)
    sheet_set = set(sheet_names)
    if {"标题台账", "关键词明细", "覆盖汇总", "Source证据"}.issubset(sheet_set):
        return "rolling_title_ledger"
    if {"关键词", "搜索指数", "搜索涨幅"}.issubset(header_set):
        return "hot_keyword_export"
    if {"产品ID", "产品名称", "搜索曝光次数", "搜索点击次数"}.issubset(header_set):
        return "product_performance_export"
    return "unknown"


def extract_dates(text: str) -> tuple[list[str], list[str]]:
    dates = []
    for y, m, d in re.findall(r"\b(20\d{2})[-_/](\d{1,2})[-_/](\d{1,2})", text):
        try:
            value = datetime(int(y), int(m), int(d)).date().isoformat()
        except ValueError:
            continue
        if value not in dates:
            dates.append(value)
    windows = []
    for match in re.findall(r"20\d{2}[-/]\d{1,2}[-/]\d{1,2}\s*[~至-]\s*20\d{2}[-/]\d{1,2}[-/]\d{1,2}", text):
        normalized = re.sub(r"\s+", "", match)
        if normalized not in windows:
            windows.append(normalized)
    return dates, windows


def read_xls(path: Path) -> list[dict]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("xlrd is required to inspect .xls files") from exc
    workbook = xlrd.open_workbook(path)
    sheets = []
    for sheet in workbook.sheets():
        preview = [sheet.row_values(index) for index in range(min(sheet.nrows, 20))]
        header_row, headers = detect_header(preview)
        sheets.append({
            "name": sheet.name,
            "rows": sheet.nrows,
            "columns": sheet.ncols,
            "header_row": header_row,
            "headers": headers,
            "preview_text": " ".join(clean(value) for row in preview[:10] for value in row if clean(value)),
        })
    return sheets


def read_xlsx(path: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to inspect .xlsx files") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        preview = []
        row_count = 0
        column_count = 0
        for row_count, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = list(row)
            column_count = max(column_count, len(values))
            if row_count <= 20:
                preview.append(values)
        header_row, headers = detect_header(preview)
        sheets.append({
            "name": sheet.title,
            "rows": row_count,
            "columns": column_count,
            "header_row": header_row,
            "headers": headers,
            "preview_text": " ".join(clean(value) for row in preview[:10] for value in row if clean(value)),
        })
    return sheets


def read_csv(path: Path) -> list[dict]:
    rows = []
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for index, row in enumerate(csv.reader(handle, delimiter=delimiter)):
            rows.append(row)
            if index >= 19:
                break
    header_row, headers = detect_header(rows)
    return [{
        "name": path.stem,
        "rows": None,
        "columns": max((len(row) for row in rows), default=0),
        "header_row": header_row,
        "headers": headers,
        "preview_text": " ".join(clean(value) for row in rows[:10] for value in row if clean(value)),
    }]


def inspect(path: Path) -> dict:
    suffix = path.suffix.lower()
    if suffix == ".xls":
        sheets = read_xls(path)
    elif suffix in {".xlsx", ".xlsm"}:
        sheets = read_xlsx(path)
    elif suffix in {".csv", ".tsv"}:
        sheets = read_csv(path)
    else:
        raise ValueError(f"Unsupported source format: {path.suffix}")
    combined_text = " ".join([path.name] + [sheet["preview_text"] for sheet in sheets])
    dates, windows = extract_dates(combined_text)
    headers = []
    for sheet in sheets:
        for value in sheet["headers"]:
            if value not in headers:
                headers.append(value)
        sheet.pop("preview_text", None)
    return {
        "original_name": path.name,
        "sha256": sha256(path),
        "size_bytes": path.stat().st_size,
        "format": suffix.lstrip("."),
        "source_type": classify([sheet["name"] for sheet in sheets], headers),
        "detected_dates": dates,
        "detected_time_windows": windows,
        "sheets": sheets,
    }


def load_registry(path: Path) -> dict:
    if not path.exists():
        return {"schema_version": 1, "sources": []}
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict) or not isinstance(data.get("sources"), list):
        raise ValueError(f"Invalid registry structure: {path}")
    return data


def atomic_write_json(path: Path, value: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temporary.replace(path)


def safe_stem(name: str) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "-", Path(name).stem).strip("-.")
    return value[:80] or "source"


def main() -> int:
    script_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser()
    parser.add_argument("files", nargs="+", type=Path)
    parser.add_argument("--registry", type=Path, default=script_root / "assets" / "source-registry.json")
    parser.add_argument("--archive-dir", type=Path, default=script_root / "assets" / "baseline-sources")
    parser.add_argument("--commit", action="store_true", help="Append new hashes and archive original bytes")
    args = parser.parse_args()

    registry = load_registry(args.registry)
    known = {record["sha256"]: record for record in registry["sources"]}
    report = []
    added_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat()

    for path in args.files:
        path = path.resolve()
        if not path.is_file():
            report.append({"file": str(path), "status": "error", "error": "not a file"})
            continue
        try:
            record = inspect(path)
        except Exception as exc:
            report.append({"file": str(path), "status": "error", "error": str(exc)})
            continue
        existing = known.get(record["sha256"])
        if existing:
            report.append({**record, "status": "duplicate", "archive_path": existing.get("archive_path")})
            continue
        status = "new"
        if args.commit:
            args.archive_dir.mkdir(parents=True, exist_ok=True)
            archive_name = f"{record['sha256'][:12]}__{safe_stem(path.name)}{path.suffix.lower()}"
            archive_path = args.archive_dir / archive_name
            if not archive_path.exists():
                shutil.copy2(path, archive_path)
            stored = {**record, "added_at": added_at, "archive_path": str(archive_path.relative_to(script_root))}
            registry["sources"].append(stored)
            known[record["sha256"]] = stored
            record = stored
            status = "registered"
        report.append({**record, "status": status})

    if args.commit:
        registry["updated_at"] = added_at
        atomic_write_json(args.registry, registry)

    result = {
        "mode": "commit" if args.commit else "inspect",
        "registry": str(args.registry),
        "files": report,
        "counts": {
            "registered": sum(item["status"] == "registered" for item in report),
            "new": sum(item["status"] == "new" for item in report),
            "duplicate": sum(item["status"] == "duplicate" for item in report),
            "error": sum(item["status"] == "error" for item in report),
        },
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 1 if result["counts"]["error"] else 0


if __name__ == "__main__":
    sys.exit(main())
